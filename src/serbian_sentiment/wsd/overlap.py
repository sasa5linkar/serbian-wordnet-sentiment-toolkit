from __future__ import annotations

import csv
import re
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from ..errors import ResourceError
from ..models import SenseDecision, TokenRecord
from ..text import normalize_lemma

WORD_RE = re.compile(r"[^\W\d_]+", re.UNICODE)


@dataclass(frozen=True)
class SenseEntry:
    id: str
    lemmas: tuple[str, ...]
    definition: str
    wordnet_pos: str = ""


@dataclass(frozen=True)
class SenseRepository:
    entries: tuple[SenseEntry, ...]
    by_lemma: dict[str, tuple[SenseEntry, ...]]


def load_sense_repository(path: str | Path) -> SenseRepository:
    source = Path(path)
    if not source.is_file():
        raise ResourceError(f"sense repository not found: {source}")
    if source.suffix.casefold() == ".xlsx":
        return _load_elexis_xlsx(source)
    return _load_csv_repository(source)


def _load_csv_repository(source: Path) -> SenseRepository:
    with source.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"ID", "Lemme", "Definicija"}
        missing = required - set(reader.fieldnames or ())
        if missing:
            raise ResourceError(
                "sense repository is missing columns: " + ", ".join(sorted(missing))
            )
        entries: list[SenseEntry] = []
        index: dict[str, list[SenseEntry]] = {}
        for row in reader:
            sense_id = (row.get("ID") or "").strip()
            lemmas = tuple(
                normalize_lemma(item)
                for item in (row.get("Lemme") or "").split(",")
                if item.strip()
            )
            if not sense_id or not lemmas:
                continue
            entry = SenseEntry(
                id=sense_id,
                lemmas=lemmas,
                definition=(row.get("Definicija") or "").strip(),
                wordnet_pos=(row.get("Vrsta") or "").strip().casefold(),
            )
            entries.append(entry)
            for lemma in lemmas:
                index.setdefault(lemma, []).append(entry)
    return _repository(entries, index)


def _load_elexis_xlsx(source: Path) -> SenseRepository:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as exc:
        raise ResourceError(
            "install the full extra to read ELEXIS XLSX resources: pip install -e .[full]"
        ) from exc
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ResourceError(f"could not read ELEXIS workbook {source}: {exc}") from exc
    sheet_name = "SrWSD-v2"
    if sheet_name not in workbook.sheetnames:
        raise ResourceError(f"ELEXIS workbook is missing sheet: {sheet_name}")
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ResourceError(f"ELEXIS sheet {sheet_name} is empty") from exc
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    positions = {name: index for index, name in enumerate(headers) if name}
    required = {"senseID", "lemma", "definition", "pos"}
    missing = required - set(positions)
    if missing:
        raise ResourceError("ELEXIS repository is missing columns: " + ", ".join(sorted(missing)))

    entries: list[SenseEntry] = []
    index: dict[str, list[SenseEntry]] = {}
    for values in rows:
        sense_id = _cell(values, positions["senseID"])
        definition = _cell(values, positions["definition"])
        primary_lemma = _cell(values, positions["lemma"])
        literal_text = _cell(values, positions["literals"]) if "literals" in positions else ""
        lemmas = tuple(
            dict.fromkeys(
                normalize_lemma(item)
                for item in (primary_lemma, *literal_text.split(","))
                if item.strip()
            )
        )
        if not sense_id or not lemmas:
            continue
        entry = SenseEntry(
            id=sense_id,
            lemmas=lemmas,
            definition=definition,
            wordnet_pos=_cell(values, positions["pos"]).casefold(),
        )
        entries.append(entry)
        for lemma in lemmas:
            index.setdefault(lemma, []).append(entry)
    workbook.close()
    if not entries:
        raise ResourceError(f"ELEXIS sheet {sheet_name} contains no usable senses")
    return _repository(entries, index)


def _cell(values: tuple[object, ...], position: int) -> str:
    if position >= len(values) or values[position] is None:
        return ""
    return str(values[position]).strip()


def _repository(
    entries: list[SenseEntry], index: dict[str, list[SenseEntry]]
) -> SenseRepository:
    return SenseRepository(tuple(entries), {key: tuple(value) for key, value in index.items()})


class OverlapSenseDisambiguator:
    """Transparent definition-overlap WSD baseline."""

    name = "definition-overlap"

    def __init__(self, repository: SenseRepository) -> None:
        self.repository = repository

    def disambiguate(
        self,
        text: str,
        tokens: Sequence[TokenRecord],
    ) -> tuple[SenseDecision, ...]:
        context = {normalize_lemma(word) for word in WORD_RE.findall(text)}
        decisions: list[SenseDecision] = []
        for token in tokens:
            candidates = self.repository.by_lemma.get(normalize_lemma(token.lemma), ())
            ranked = sorted(
                candidates,
                key=lambda sense: (
                    -len(
                        context
                        & {normalize_lemma(word) for word in WORD_RE.findall(sense.definition)}
                    ),
                    sense.id,
                ),
            )
            selected = ranked[0].id if ranked else None
            decisions.append(
                SenseDecision(
                    token=token,
                    selected_sense_id=selected,
                    candidates=tuple(item.id for item in ranked),
                    backend=self.name,
                    explanation=(
                        f"Selected by definition overlap from {len(ranked)} candidate(s)."
                        if ranked
                        else "No candidate sense was found for the normalized lemma."
                    ),
                )
            )
        return tuple(decisions)
